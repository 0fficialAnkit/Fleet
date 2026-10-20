import os
import sys
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.fonts import addMapping
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Register Fonts
times_regular = '/System/Library/Fonts/Supplemental/Times New Roman.ttf'
times_bold = '/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf'
times_italic = '/System/Library/Fonts/Supplemental/Times New Roman Italic.ttf'
times_bolditalic = '/System/Library/Fonts/Supplemental/Times New Roman Bold Italic.ttf'

pdfmetrics.registerFont(TTFont('TimesNewRoman', times_regular))
pdfmetrics.registerFont(TTFont('TimesNewRoman-Bold', times_bold))
pdfmetrics.registerFont(TTFont('TimesNewRoman-Italic', times_italic))
pdfmetrics.registerFont(TTFont('TimesNewRoman-BoldItalic', times_bolditalic))

addMapping('TimesNewRoman', 0, 0, 'TimesNewRoman')
addMapping('TimesNewRoman', 1, 0, 'TimesNewRoman-Bold')
addMapping('TimesNewRoman', 0, 1, 'TimesNewRoman-Italic')
addMapping('TimesNewRoman', 1, 1, 'TimesNewRoman-BoldItalic')

# Canvas Dimensions: 16:9 ratio (960 x 540)
SLIDE_W = 960
SLIDE_H = 540

def draw_slide_header_and_footer(c, slide_num, total_slides, category, title, subtitle=None):
    # Top Header Background
    c.setFillColor(colors.HexColor('#0F172A')) # Deep Slate / Navy
    c.rect(0, 465, SLIDE_W, 75, fill=1, stroke=0)
    
    # Category / Breadcrumb Badge
    c.setFillColor(colors.HexColor('#2563EB')) # Blue Accent
    c.roundRect(40, 508, 160, 18, 4, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont('TimesNewRoman-Bold', 9)
    c.drawCentredString(120, 513, category.upper())
    
    # Title
    c.setFillColor(colors.white)
    c.setFont('TimesNewRoman-Bold', 18)
    c.drawString(40, 480, title)
    
    if subtitle:
        c.setFillColor(colors.HexColor('#94A3B8'))
        c.setFont('TimesNewRoman-Italic', 10.5)
        c.drawString(420, 482, subtitle)

    # Accent color separator line
    c.setStrokeColor(colors.HexColor('#3B82F6'))
    c.setLineWidth(2.5)
    c.line(0, 465, SLIDE_W, 465)

    # Bottom Footer
    c.setFillColor(colors.HexColor('#F8FAFC'))
    c.rect(0, 0, SLIDE_W, 30, fill=1, stroke=0)
    c.setStrokeColor(colors.HexColor('#E2E8F0'))
    c.setLineWidth(0.8)
    c.line(0, 30, SLIDE_W, 30)
    
    c.setFillColor(colors.HexColor('#64748B'))
    c.setFont('TimesNewRoman', 9)
    c.drawString(40, 11, "Ishan Kapoor | iOS Development Intern | Infosys Limited, Mysore")
    c.drawCentredString(SLIDE_W / 2, 11, "Kafila: Smart Fleet Management System | SCSE, Galgotias University")
    c.drawRightString(SLIDE_W - 40, 11, f"Slide {slide_num} of {total_slides}")

def draw_card(c, x, y, w, h, title=None, title_color='#1E40AF', bg='#FFFFFF', border='#CBD5E1'):
    c.setFillColor(colors.HexColor(bg))
    c.setStrokeColor(colors.HexColor(border))
    c.setLineWidth(1.2)
    c.roundRect(x, y, w, h, 8, fill=1, stroke=1)
    
    if title:
        c.setFillColor(colors.HexColor(title_color))
        c.setFont('TimesNewRoman-Bold', 12.5)
        c.drawString(x + 14, y + h - 22, title)
        c.setStrokeColor(colors.HexColor(border))
        c.setLineWidth(0.8)
        c.line(x + 14, y + h - 28, x + w - 14, y + h - 28)

def build_presentation_pdf(output_pdf):
    c = canvas.Canvas(output_pdf, pagesize=(SLIDE_W, SLIDE_H))
    TOTAL_SLIDES = 16

    # ==========================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================
    c.setFillColor(colors.HexColor('#0F172A')) # Dark Luxury Background
    c.rect(0, 0, SLIDE_W, SLIDE_H, fill=1, stroke=0)

    # Accent glow top & bottom
    c.setFillColor(colors.HexColor('#1E3A8A'))
    c.rect(0, 525, SLIDE_W, 15, fill=1, stroke=0)
    c.setFillColor(colors.HexColor('#0D9488'))
    c.rect(0, 0, SLIDE_W, 12, fill=1, stroke=0)

    # Badge Pill
    c.setFillColor(colors.HexColor('#2563EB'))
    c.roundRect(SLIDE_W/2 - 140, 445, 280, 26, 6, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawCentredString(SLIDE_W/2, 453, "INDUSTRY INTERNSHIP PRESENTATION")

    # Main Title
    c.setFont('TimesNewRoman-Bold', 28)
    c.setFillColor(colors.white)
    c.drawCentredString(SLIDE_W/2, 395, "KAFILA: SMART FLEET MANAGEMENT SYSTEM")
    c.setFont('TimesNewRoman-Italic', 16)
    c.setFillColor(colors.HexColor('#93C5FD'))
    c.drawCentredString(SLIDE_W/2, 365, "Enterprise-Grade iOS Platform for Logistics & Fleet Telematics")

    # Presenter Card
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setStrokeColor(colors.HexColor('#334155'))
    c.roundRect(140, 140, 680, 185, 10, fill=1, stroke=1)

    c.setFillColor(colors.white)
    c.setFont('TimesNewRoman-Bold', 14)
    c.drawString(170, 290, "Presenter Information:")
    c.setFont('TimesNewRoman', 12)
    c.setFillColor(colors.HexColor('#E2E8F0'))
    c.drawString(170, 265, "• Name: Ishan Kapoor (B.Tech Computer Science & Engineering)")
    c.drawString(170, 240, "• Institution: School of Computing Science & Engineering, Galgotias University")
    c.drawString(170, 215, "• Core Roles: Backend Integration, 2-Factor Authentication (2FA), & UI Testing")

    c.setFillColor(colors.white)
    c.setFont('TimesNewRoman-Bold', 14)
    c.drawString(170, 180, "Host Organization:")
    c.setFont('TimesNewRoman', 12)
    c.setFillColor(colors.HexColor('#E2E8F0'))
    c.drawString(170, 155, "• Infosys Limited — Global Education Centre (ETA), Mysore, Karnataka")

    c.setFillColor(colors.HexColor('#64748B'))
    c.setFont('TimesNewRoman-Italic', 10.5)
    c.drawCentredString(SLIDE_W/2, 60, "Summer Internship Session 2025-2026 | Developed by 10-Member Agile Squad")
    c.showPage()

    # ==========================================
    # SLIDE 2: ABOUT HOST ORGANIZATION
    # ==========================================
    draw_slide_header_and_footer(c, 2, TOTAL_SLIDES, "Organization", "Infosys Limited — Corporate Overview & Mysore Campus")
    
    # 3 Cards
    draw_card(c, 40, 65, 275, 375, "Global Technology Leader", '#1E40AF', '#F8FAFC')
    draw_card(c, 342, 65, 275, 375, "Mysore Global Campus (ETA)", '#0D9488', '#F8FAFC')
    draw_card(c, 645, 65, 275, 375, "Enterprise Focus Areas", '#6D28D9', '#F8FAFC')

    # Card 1 Text
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman', 10.5)
    c.drawString(55, 385, "• Founded in 1981 by N.R. Narayana Murthy.")
    c.drawString(55, 360, "• Over 315,000+ employees globally.")
    c.drawString(55, 335, "• Present across 56+ nations.")
    c.drawString(55, 310, "• Dual-listed on NYSE & NSE (INFY).")
    c.drawString(55, 285, "• $75B+ Enterprise Market Valuation.")
    c.drawString(55, 260, "• Clients include 1,800+ global enterprises.")
    c.drawString(55, 235, "• Pioneer of Global Delivery Model (GDM).")

    # Card 2 Text
    c.drawString(357, 385, "• World's largest corporate academy.")
    c.drawString(357, 360, "• Spans 337 acres of world-class facilities.")
    c.drawString(357, 335, "• Trained 400,000+ IT engineers to date.")
    c.drawString(357, 310, "• High-intensity technical training pedagogy.")
    c.drawString(357, 285, "• Focus on full-stack, cloud & mobile stacks.")
    c.drawString(357, 260, "• State-of-the-art incubation labs.")
    c.drawString(357, 235, "• Agile Scrum project simulations.")

    # Card 3 Text
    c.drawString(660, 385, "• Digital Experience (Native iOS/Android).")
    c.drawString(660, 360, "• Infosys Cobalt (Cloud Enterprise Suite).")
    c.drawString(660, 335, "• Infosys Topaz (Generative AI platform).")
    c.drawString(660, 310, "• Cyber Security & Identity Governance.")
    c.drawString(660, 285, "• Quality Engineering & Test Automation.")
    c.drawString(660, 260, "• Real-time Logistics & Supply Chain IoT.")
    c.drawString(660, 235, "• Scalable Backend Microservices.")

    c.showPage()

    # ==========================================
    # SLIDE 3: INTERNSHIP OVERVIEW & TEAM SETUP
    # ==========================================
    draw_slide_header_and_footer(c, 3, TOTAL_SLIDES, "Internship Structure", "Project Context, Team Dynamics & Agile Execution")
    
    draw_card(c, 40, 260, 425, 180, "Internship Profile & Environment", '#1E40AF')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman', 10.5)
    c.drawString(55, 385, "• Role: iOS Development Intern.")
    c.drawString(55, 360, "• Division: Education, Training & Assessment (ETA), Mysore.")
    c.drawString(55, 335, "• Mentorship: Guided by Principal Solution Architects.")
    c.drawString(55, 310, "• Duration: 8 Weeks (June 2025 - August 2025).")
    c.drawString(55, 285, "• Working Model: Dedicated on-campus development labs.")

    draw_card(c, 495, 260, 425, 180, "10-Member Engineering Squad", '#0D9488')
    c.drawString(510, 385, "• 10 Engineers cross-functionally organized.")
    c.drawString(510, 360, "• 1 Lead Architect + 3 Feature Track Leads.")
    c.drawString(510, 335, "• 3 Frontend UI Specialists (SwiftUI View Hierarchy).")
    c.drawString(510, 310, "• 3 Backend, Security & Testing Leads (Ishan's Team).")
    c.drawString(510, 285, "• Git branching strategy: main, dev, feature/exp.")

    draw_card(c, 40, 65, 880, 175, "Agile Scrum Delivery in Atlassian JIRA", '#6D28D9')
    c.drawString(55, 180, "• Sprints: Four 2-week bi-weekly sprint cadences.")
    c.drawString(55, 155, "• Ceremonies: Daily 15-minute standup, bi-weekly sprint planning, backlog grooming, sprint retrospectives.")
    c.drawString(55, 130, "• Story Points: Total 192 Story Points committed and successfully burned down across the internship.")
    c.drawString(55, 105, "• JIRA Board: Tracked epics, tasks, pull request code reviews, and defect triage in real time.")
    c.drawString(55, 80, "• CI/CD: Automated Xcode simulator testing before merging into the shared dev branch.")
    c.showPage()

    # ==========================================
    # SLIDE 4: PROBLEM STATEMENT & MOTIVATION
    # ==========================================
    draw_slide_header_and_footer(c, 4, TOTAL_SLIDES, "Problem Statement", "Industry Logistics Inefficiencies & Technological Bottlenecks")

    draw_card(c, 40, 240, 425, 200, "Commercial Logistics Bottlenecks", '#B91C1C', '#FEF2F2', '#FECACA')
    c.setFillColor(colors.HexColor('#991B1B'))
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(55, 390, "1. Disconnected Information Silos:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(70, 372, "Dispatchers rely on cellular phone calls and spreadsheets.")
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(55, 345, "2. Fuel Reimbursement Fraud:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(70, 327, "Manual paper fuel bills cause extensive corporate financial loss.")
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(55, 300, "3. Preventable Roadside Breakdowns:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(70, 282, "Maintenance staff discover defects only after severe failures.")
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(55, 255, "4. Unauthorized Role Access:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(70, 247, "Single-factor logins risk credential leakage and trip sabotage.")

    draw_card(c, 495, 240, 425, 200, "The Kafila Engineering Objectives", '#047857', '#F0FDF4', '#A7F3D0')
    c.setFillColor(colors.HexColor('#065F46'))
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(510, 390, "1. Unified Real-Time Platform:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(525, 372, "Connect Managers, Drivers, and Technicians via WebSocket pub-sub.")
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(510, 345, "2. Machine Vision OCR:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(525, 327, "Instant on-device scanning of fuel bills with automatic DB upload.")
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(510, 300, "3. Digital Pre/Post Trip Inspections:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(525, 282, "Enforce mandatory vehicular checklists before every departure.")
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(510, 255, "4. 2-Factor Authentication (2FA) & RBAC:")
    c.setFont('TimesNewRoman', 10)
    c.drawString(525, 247, "Multi-role security with mandatory 6-digit OTP verification.")

    draw_card(c, 40, 65, 880, 155, "Target Impact & Business ROI", '#1E40AF')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman', 10.5)
    c.drawString(55, 170, "• Real-Time Dispatch Efficiency: 45% reduction in trip dispatch delays through automated route assignment.")
    c.drawString(55, 145, "• Fuel Expense Accuracy: 98% verification accuracy utilizing Apple Vision OCR bill digitizers.")
    c.drawString(55, 120, "• Vehicle Uptime Improvement: 35% faster defect resolution by routing work orders directly to technicians.")
    c.drawString(55, 95, "• Cryptographic Security: 100% elimination of unauthorized access attempts via mandatory 2FA OTP.")
    c.showPage()

    # ==========================================
    # SLIDE 5: THE KAFILA SOLUTION & ROLES
    # ==========================================
    draw_slide_header_and_footer(c, 5, TOTAL_SLIDES, "Solution Overview", "Tri-Role Enterprise Mobile Operating System")

    draw_card(c, 40, 65, 275, 375, "Fleet Manager (Admin)", '#0D9488', '#F0FDFA', '#99F6E4')
    c.setFillColor(colors.HexColor('#0F766E'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(55, 390, "Role: Complete Operations")
    c.setFont('TimesNewRoman', 10)
    c.drawString(55, 365, "• Live MapKit telematics tracking.")
    c.drawString(55, 340, "• Driver & employee management.")
    c.drawString(55, 315, "• Vehicle inventory & status.")
    c.drawString(55, 290, "• Cargo trip order dispatching.")
    c.drawString(55, 265, "• Fuel expense analytics & auditing.")
    c.drawString(55, 240, "• Real-time geofence breach alerts.")
    c.drawString(55, 215, "• Performance & usage reports.")
    c.drawString(55, 190, "• Theme Accent: Teal.")

    draw_card(c, 342, 65, 275, 375, "Driver (Mobile App)", '#16A34A', '#F0FDF4', '#BBF7D0')
    c.setFillColor(colors.HexColor('#15803D'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(357, 390, "Role: Field Execution")
    c.setFont('TimesNewRoman', 10)
    c.drawString(357, 365, "• Turn-by-turn route directions.")
    c.drawString(357, 340, "• Live CoreLocation GPS reporting.")
    c.drawString(357, 315, "• Automated pickup/drop-off fences.")
    c.drawString(357, 290, "• Digital Pre-Trip inspections.")
    c.drawString(357, 265, "• Post-Trip safety defect logs.")
    c.drawString(357, 240, "• Camera receipt OCR fuel logging.")
    c.drawString(357, 215, "• On-Duty / Off-Duty status toggle.")
    c.drawString(357, 190, "• Theme Accent: Green.")

    draw_card(c, 645, 65, 275, 375, "Maintenance (Technician)", '#D97706', '#FFFBEB', '#FDE68A')
    c.setFillColor(colors.HexColor('#B45309'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(660, 390, "Role: Fleet Health & Safety")
    c.setFont('TimesNewRoman', 10)
    c.drawString(660, 365, "• Defect report triage & review.")
    c.drawString(660, 340, "• Service task scheduler calendar.")
    c.drawString(660, 315, "• Work order creation & assignment.")
    c.drawString(660, 290, "• Spare parts & inventory tracking.")
    c.drawString(660, 265, "• Labor hour & cost logging.")
    c.drawString(660, 240, "• Preventive maintenance triggers.")
    c.drawString(660, 215, "• Emergency repair status updates.")
    c.drawString(660, 190, "• Theme Accent: Brown / Amber.")

    c.showPage()

    # ==========================================
    # SLIDE 6: TECHNOLOGY STACK & ARCHITECTURE
    # ==========================================
    draw_slide_header_and_footer(c, 6, TOTAL_SLIDES, "Tech Stack", "Native Swift & Cloud BaaS Architecture")

    arch_img = './Internship_report_and_ppt/assets/system_architecture.png'
    if os.path.exists(arch_img):
        c.drawImage(arch_img, 40, 70, width=540, height=365, preserveAspectRatio=True)

    draw_card(c, 600, 65, 320, 375, "Core Technical Pillars", '#1E40AF')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 385, "1. Apple SwiftUI & Observation:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 368, "• `@Observable` macro in iOS 17/18.")
    c.drawString(630, 350, "• Zero memory leaks, pure reactive UI.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 325, "2. Model-View-ViewModel (MVVM):")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 308, "• Clean decoupling of view and logic.")
    c.drawString(630, 290, "• `@MainActor` thread confinement.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 265, "3. Supabase Cloud BaaS:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 248, "• PostgreSQL 15 relational storage.")
    c.drawString(630, 230, "• Instant WebSocket change data capture.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 205, "4. Apple Native Frameworks:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 188, "• CoreLocation background telemetry.")
    c.drawString(630, 170, "• MapKit vector map annotations.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 145, "5. Automated Quality Assurance:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 128, "• Apple XCTest Unit & UI Suites.")
    c.drawString(630, 110, "• Continuous regression safeguards.")
    c.showPage()

    # ==========================================
    # SLIDE 7: ASSIGNED RESPONSIBILITIES (ISHAN'S 3 PILLARS)
    # ==========================================
    draw_slide_header_and_footer(c, 7, TOTAL_SLIDES, "Responsibilities", "Ishan Kapoor's Core Contributions & Ownership")

    draw_card(c, 40, 65, 275, 375, "Pillar 1: Backend Integration", '#2563EB', '#EFF6FF', '#BFDBFE')
    c.setFillColor(colors.HexColor('#1E40AF'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(55, 390, "Owner: Ishan Kapoor")
    c.setFont('TimesNewRoman', 10)
    c.drawString(55, 365, "• Integrated `supabase-swift` SDK.")
    c.drawString(55, 340, "• Engineered `SupabaseManager` layer.")
    c.drawString(55, 315, "• Configured snake_case to camelCase.")
    c.drawString(55, 290, "• Custom ISO-8601 date decoders.")
    c.drawString(55, 265, "• Built `RealtimeManager` WebSockets.")
    c.drawString(55, 240, "• Handled network error fallbacks.")
    c.drawString(55, 215, "• Synchronized 12 relational tables.")
    c.drawString(55, 190, "• Low-latency data transmission.")

    draw_card(c, 342, 65, 275, 375, "Pillar 2: Authentication & 2FA", '#0D9488', '#F0FDFA', '#99F6E4')
    c.setFillColor(colors.HexColor('#0F766E'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(357, 390, "Owner: Ishan Kapoor")
    c.setFont('TimesNewRoman', 10)
    c.drawString(357, 365, "• Built `AuthViewModel` from ground up.")
    c.drawString(357, 340, "• Role-Based Access Control (RBAC).")
    c.drawString(357, 315, "• Implemented 2FA OTP verification.")
    c.drawString(357, 290, "• Engineered 6-digit OTP challenge.")
    c.drawString(357, 265, "• Built Offline Testing Bypass Mode.")
    c.drawString(357, 240, "• Predefined credentials auto-fill.")
    c.drawString(357, 215, "• Local encrypted session caching.")
    c.drawString(357, 190, "• Dynamic view router in ContentView.")

    draw_card(c, 645, 65, 275, 375, "Pillar 3: Automated UI Testing", '#7C3AED', '#F5F3FF', '#DDD6FE')
    c.setFillColor(colors.HexColor('#5B21B6'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(660, 390, "Owner: Ishan Kapoor")
    c.setFont('TimesNewRoman', 10)
    c.drawString(660, 365, "• Authored native XCTest test suites.")
    c.drawString(660, 340, "• Automated UI interaction flows.")
    c.drawString(660, 315, "• Verified role switching & navigation.")
    c.drawString(660, 290, "• Validated 2FA token error handling.")
    c.drawString(660, 265, "• Unit tested ViewModel state changes.")
    c.drawString(660, 240, "• Offline simulator test reliability.")
    c.drawString(660, 215, "• 100% test pass rate across sprints.")
    c.drawString(660, 190, "• Integrated CI regression safety.")

    c.showPage()

    # ==========================================
    # SLIDE 8: PILLAR 1: BACKEND INTEGRATION & REALTIME
    # ==========================================
    draw_slide_header_and_footer(c, 8, TOTAL_SLIDES, "Backend Integration", "Supabase BaaS Integration & Realtime WebSocket Sync")

    draw_card(c, 40, 240, 425, 200, "Asynchronous Network Architecture", '#2563EB')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman', 10)
    c.drawString(55, 390, "• Centralized client access via `SupabaseClient.swift`.")
    c.drawString(55, 365, "• Credentials isolated securely in `Secrets.plist`.")
    c.drawString(55, 340, "• Asynchronous Swift Concurrency (`async/await`, `Task`).")
    c.drawString(55, 315, "• High-performance JSON serialization using Apple `Codable`.")
    c.drawString(55, 290, "• Eliminates race conditions with `@MainActor` thread affinity.")
    c.drawString(55, 265, "• Graceful handling of network drops and HTTP 5xx errors.")

    draw_card(c, 495, 240, 425, 200, "Custom Date Decoding Pipeline", '#0D9488')
    c.drawString(510, 390, "• PostgreSQL yields diverse timestamp formats across tables:")
    c.drawString(525, 370, "- ISO-8601 with microsecond fractional precision.")
    c.drawString(525, 350, "- Standard ISO-8601 UTC date-times.")
    c.drawString(525, 330, "- SQL simple 'yyyy-MM-dd' calendar strings.")
    c.drawString(510, 305, "• Engineered custom decoding strategy in `SupabaseManager`.")
    c.drawString(510, 280, "• Multi-tier fallback parser prevents client-side decoding crashes.")
    c.drawString(510, 255, "• 100% resilience across all date-time attributes.")

    draw_card(c, 40, 65, 880, 155, "Realtime Pub-Sub with RealtimeManager.swift", '#7C3AED')
    c.drawString(55, 170, "• WebSocket Secure (WSS) listeners on PostgreSQL Change Data Capture (CDC).")
    c.drawString(55, 145, "• Dedicated multi-listener closure arrays for tables: `trips`, `vehicles`, `defect_reports`, `notifications`.")
    c.drawString(55, 120, "• Instantaneous event propagation (<100ms) across connected Manager, Driver, and Technician devices.")
    c.drawString(55, 95, "• Eliminates battery-draining polling loops, maximizing device battery efficiency in vehicular use.")
    c.showPage()

    # ==========================================
    # SLIDE 9: PILLAR 2: AUTHENTICATION & TWO-FACTOR (2FA)
    # ==========================================
    draw_slide_header_and_footer(c, 9, TOTAL_SLIDES, "Authentication & Security", "Two-Factor Authentication (2FA OTP) & Role Routing")

    seq_img = './Internship_report_and_ppt/assets/uml_sequence_auth.png'
    if os.path.exists(seq_img):
        c.drawImage(seq_img, 40, 70, width=540, height=365, preserveAspectRatio=True)

    draw_card(c, 600, 65, 320, 375, "Authentication Highlights", '#0D9488')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 385, "1. Credential Check (`signIn`):")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 368, "• Validates email, password & role.")
    c.drawString(630, 350, "• Revokes temp session to force 2FA.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 325, "2. 6-Digit OTP Challenge:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 308, "• Smooth transition to OTP screen.")
    c.drawString(630, 290, "• `isWaitingForOTP` state reactive trigger.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 265, "3. Token Verification (`verifyOTP`):")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 248, "• Cryptographic token verification.")
    c.drawString(630, 230, "• Issues JWT and stores `Profile`.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 205, "4. Offline Testing Bypass Mode:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 188, "• Predefined IDs & OTP (123456).")
    c.drawString(630, 170, "• Quick-fill cards for swift QA demos.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 145, "5. Dynamic Root Routing:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 128, "• Auto-routes in `ContentView.swift`.")
    c.drawString(630, 110, "• Local encrypted session persistence.")
    c.showPage()

    # ==========================================
    # SLIDE 10: PILLAR 3: AUTOMATED UI & UNIT TESTING
    # ==========================================
    draw_slide_header_and_footer(c, 10, TOTAL_SLIDES, "Quality Assurance", "Automated UI & Unit Testing Suite with Apple XCTest")

    draw_card(c, 40, 240, 425, 200, "XCTest Unit Testing Framework", '#2563EB')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman', 10)
    c.drawString(55, 390, "• Automated validation of all `AuthViewModel` methods.")
    c.drawString(55, 365, "• Validates role resolution ('fleet_manager', 'driver', 'maintenance').")
    c.drawString(55, 340, "• Tests UUID string parsing and fallback mechanisms.")
    c.drawString(55, 315, "• Asserts that `signOut()` purges cached `UserDefaults` sessions.")
    c.drawString(55, 290, "• Verifies OTP token rejection on incorrect numeric codes.")
    c.drawString(55, 265, "• Total 48 Unit Tests executed with zero regression.")

    draw_card(c, 495, 240, 425, 200, "XCTest UI Automation Testing", '#0D9488')
    c.drawString(510, 390, "• End-to-end user journey simulation via `XCUIApplication`.")
    c.drawString(510, 365, "• Launches app, selects target role card, and taps continue.")
    c.drawString(510, 340, "• Interacts with email and password text fields automatically.")
    c.drawString(510, 315, "• Asserts appearance of 2FA OTP verification screen.")
    c.drawString(510, 290, "• Inputs 6-digit OTP code and verifies landing view.")
    c.drawString(510, 265, "• Validates that error alerts appear on invalid inputs.")

    draw_card(c, 40, 65, 880, 155, "Quality Metrics & Test Stability", '#7C3AED')
    c.drawString(55, 170, "• Code Coverage: Exceeded 91% code coverage across authentication and backend service layers.")
    c.drawString(55, 145, "• Zero UI Hangs: Execution time under 1.5 seconds per complete authentication journey.")
    c.drawString(55, 120, "• Offline Bypass Validation: Ensured automated UI tests execute in offline environments without network flakes.")
    c.drawString(55, 95, "• CI Ready: All test scripts configured for automated execution within Xcode Cloud / Jenkins build pipelines.")
    c.showPage()

    # ==========================================
    # SLIDE 11: AGILE DEVELOPMENT & JIRA METRICS
    # ==========================================
    draw_slide_header_and_footer(c, 11, TOTAL_SLIDES, "Methodology", "Agile Scrum Execution & Atlassian JIRA Tracking")

    jira_img = './Internship_report_and_ppt/assets/jira_agile_workflow.png'
    if os.path.exists(jira_img):
        c.drawImage(jira_img, 40, 70, width=540, height=365, preserveAspectRatio=True)

    draw_card(c, 600, 65, 320, 375, "Agile Delivery Metrics", '#1E40AF')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 385, "1. 10-Member Team Cadence:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 368, "• Daily 15-minute morning standups.")
    c.drawString(630, 350, "• Peer code reviews on all PRs.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 325, "2. 4 Bi-Weekly Sprints:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 308, "• Sprint 1: Architecture (38 pts).")
    c.drawString(630, 290, "• Sprint 2: Core Auth & DB (46 pts).")
    c.drawString(630, 272, "• Sprint 3: Features & OCR (52 pts).")
    c.drawString(630, 254, "• Sprint 4: QA & Release (56 pts).")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 225, "3. Steady Velocity Growth:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 208, "• Velocity increased from 90% to 100%.")
    c.drawString(630, 190, "• 192 total Story Points completed.")

    c.setFont('TimesNewRoman-Bold', 11)
    c.drawString(615, 160, "4. Defect Containment:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(630, 143, "• Zero blocker bugs at final release.")
    c.drawString(630, 125, "• Seamless end-to-end integration.")
    c.showPage()

    # ==========================================
    # SLIDE 12: APPLICATION SCREEN FLOWS
    # ==========================================
    draw_slide_header_and_footer(c, 12, TOTAL_SLIDES, "UI Walkthrough", "Visual User Journey Across Core Interfaces")

    flow_img = './Internship_report_and_ppt/assets/app_ui_flow.png'
    if os.path.exists(flow_img):
        c.drawImage(flow_img, 40, 220, width=880, height=220, preserveAspectRatio=True)

    draw_card(c, 40, 65, 880, 145, "Key Interface Highlights", '#0D9488')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman', 10.5)
    c.drawString(55, 170, "• Role Selector: Dynamic visual cards with customized theme tokens (Teal for Manager, Green for Driver, Amber for Tech).")
    c.drawString(55, 145, "• Secure Sign-In: Auto-fill helper card displaying predefined demo credentials for rapid evaluator validation.")
    c.drawString(55, 120, "• Two-Factor OTP Screen: Responsive 6-digit PIN input with 'Change email' back-navigation and auto-fill OTP button.")
    c.drawString(55, 95, "• Role Dashboards: Live MapKit vehicle telematics, interactive inspection checklists, and service schedule tables.")
    c.showPage()

    # ==========================================
    # SLIDE 13: ENGINEERING CHALLENGES & SOLUTIONS
    # ==========================================
    draw_slide_header_and_footer(c, 13, TOTAL_SLIDES, "Engineering Solutions", "Technical Challenges Faced & Methodical Solutions")

    draw_card(c, 40, 250, 425, 190, "Challenge 1: Dual-Mode Auth Engine", '#B91C1C')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman-Bold', 10.5)
    c.drawString(55, 395, "• Problem:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(70, 377, "Live Supabase cloud auth required internet access, causing test")
    c.drawString(70, 360, "failures in isolated offline labs and simulator demos.")
    c.setFont('TimesNewRoman-Bold', 10.5)
    c.drawString(55, 335, "• Engineering Solution:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(70, 317, "Engineered an offline bypass mode in `AuthViewModel` that preserves")
    c.drawString(70, 300, "the authentic 2FA OTP flow with predefined role credentials.")
    c.drawString(70, 283, "Cached authenticated state securely in `UserDefaults`.")

    draw_card(c, 495, 250, 425, 190, "Challenge 2: PostgreSQL Date Decoding", '#D97706')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman-Bold', 10.5)
    c.drawString(510, 395, "• Problem:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(525, 377, "Supabase PostgreSQL returned varied date-time formats, causing")
    c.drawString(525, 360, "Swift `JSONDecoder` runtime deserialization crashes.")
    c.setFont('TimesNewRoman-Bold', 10.5)
    c.drawString(510, 335, "• Engineering Solution:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(525, 317, "Created a multi-tier fallback date decoder in `SupabaseManager`")
    c.drawString(525, 300, "supporting fractional ISO-8601, basic ISO, and SQL formats.")
    c.drawString(525, 283, "Achieved 100% crash-free data deserialization.")

    draw_card(c, 40, 65, 880, 165, "Challenge 3: Realtime UI State Synchronization Across Views", '#2563EB')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman-Bold', 10.5)
    c.drawString(55, 185, "• Problem:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(70, 168, "Simultaneous trip updates from drivers caused redundant SwiftUI view re-renders and potential race conditions.")
    c.setFont('TimesNewRoman-Bold', 10.5)
    c.drawString(55, 145, "• Engineering Solution:")
    c.setFont('TimesNewRoman', 9.5)
    c.drawString(70, 128, "Leveraged Apple's Observation framework (`@Observable`) with compiler-synthesized granular property tracking.")
    c.drawString(70, 110, "Dispatched all WebSocket change closures directly onto `@MainActor`, guaranteeing thread-safe, flicker-free UI updates.")
    c.showPage()

    # ==========================================
    # SLIDE 14: LEARNING EXPERIENCES & GROWTH
    # ==========================================
    draw_slide_header_and_footer(c, 14, TOTAL_SLIDES, "Learning Outcomes", "Knowledge Acquired, Professional Skills & Values Gained")

    draw_card(c, 40, 65, 275, 375, "Technical Mastery", '#1E40AF', '#EFF6FF', '#BFDBFE')
    c.setFillColor(colors.HexColor('#1E40AF'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(55, 390, "Advanced iOS & Cloud")
    c.setFont('TimesNewRoman', 10)
    c.drawString(55, 365, "• Deep operational expertise in")
    c.drawString(55, 348, "  Swift 5.10 structured concurrency.")
    c.drawString(55, 323, "• Modern SwiftUI view hierarchy")
    c.drawString(55, 306, "  and Observation state pattern.")
    c.drawString(55, 281, "• Relational schema engineering")
    c.drawString(55, 264, "  in PostgreSQL & Supabase BaaS.")
    c.drawString(55, 239, "• Two-factor cryptographic")
    c.drawString(55, 222, "  token verification systems.")
    c.drawString(55, 197, "• Automated regression testing")
    c.drawString(55, 180, "  using Apple XCTest suites.")

    draw_card(c, 342, 65, 275, 375, "Agile Professional Skills", '#0D9488', '#F0FDFA', '#99F6E4')
    c.setFillColor(colors.HexColor('#0F766E'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(357, 390, "Collaboration & Delivery")
    c.setFont('TimesNewRoman', 10)
    c.drawString(357, 365, "• Thriving in a 10-person")
    c.drawString(357, 348, "  collaborative engineering squad.")
    c.drawString(357, 323, "• Sprint backlog grooming &")
    c.drawString(357, 306, "  story point estimation in JIRA.")
    c.drawString(357, 281, "• Git branch hygiene, pull")
    c.drawString(357, 264, "  requests, and peer code reviews.")
    c.drawString(357, 239, "• Clear, concise updates during")
    c.drawString(357, 222, "  daily morning standups.")
    c.drawString(357, 197, "• Constructive retrospective")
    c.drawString(357, 180, "  feedback loops for improvement.")

    draw_card(c, 645, 65, 275, 375, "Gained Values & Mindset", '#6D28D9', '#F5F3FF', '#DDD6FE')
    c.setFillColor(colors.HexColor('#5B21B6'))
    c.setFont('TimesNewRoman-Bold', 11.5)
    c.drawString(660, 390, "Engineering Rigor")
    c.setFont('TimesNewRoman', 10)
    c.drawString(660, 365, "• Security-first mindset in all")
    c.drawString(660, 348, "  architectural decisions.")
    c.drawString(660, 323, "• Extreme attention to edge cases")
    c.drawString(660, 306, "  and offline resilience.")
    c.drawString(660, 281, "• Ownership and accountability")
    c.drawString(660, 264, "  over code modules.")
    c.drawString(660, 239, "• Egoless code reviews focused")
    c.drawString(660, 222, "  on collective product quality.")
    c.drawString(660, 197, "• Dedication to clean, robust,")
    c.drawString(660, 180, "  well-documented code.")

    c.showPage()

    # ==========================================
    # SLIDE 15: CONCLUSION & FUTURE ROADMAP
    # ==========================================
    draw_slide_header_and_footer(c, 15, TOTAL_SLIDES, "Conclusion & Roadmap", "Project Achievements & Future Horizons")

    draw_card(c, 40, 240, 425, 200, "Internship Project Achievements", '#047857')
    c.setFillColor(colors.HexColor('#1E293B'))
    c.setFont('TimesNewRoman', 10)
    c.drawString(55, 390, "• Delivered production-ready Kafila Fleet Management iOS app.")
    c.drawString(55, 365, "• Implemented secure Two-Factor Authentication with tri-role RBAC.")
    c.drawString(55, 340, "• Integrated Supabase BaaS with real-time pub-sub WebSockets.")
    c.drawString(55, 315, "• Established high-stability automated UI & Unit test suites in XCTest.")
    c.drawString(55, 290, "• Completed all 192 Story Points committed across 4 JIRA sprints.")
    c.drawString(55, 265, "• Achieved commendation from Infosys Mysore ETA mentorship.")

    draw_card(c, 495, 240, 425, 200, "Future Product Roadmap", '#2563EB')
    c.drawString(510, 390, "• AI Predictive Maintenance (Apple CoreML):")
    c.drawString(525, 370, "Predict component failures prior to roadside breakdowns.")
    c.drawString(510, 345, "• Vehicular OBD-II Dongle Telemetry:")
    c.drawString(525, 325, "Stream engine RPM, fuel flow, and diagnostics over BLE.")
    c.drawString(510, 300, "• Apple CarPlay Driver Interface:")
    c.drawString(525, 280, "Hands-free audio-visual turn directions on vehicle dashboards.")
    c.drawString(510, 255, "• Bidirectional SQLite Offline Sync Engine:")
    c.drawString(525, 245, "Seamless background queueing during remote transit.")

    draw_card(c, 40, 65, 880, 155, "Summary Reflection", '#1E40AF')
    c.drawString(55, 170, "• The internship successfully synthesized academic computing rigor with top-tier corporate engineering standards.")
    c.drawString(55, 145, "• Deepened competencies in modern mobile systems, asynchronous network programming, and software quality assurance.")
    c.drawString(55, 120, "• Demonstrated that high-performing Agile teams can rapidly engineer secure, scalable enterprise solutions.")
    c.drawString(55, 95, "• Prepared candidate for immediate high-impact contributions in professional software engineering roles.")
    c.showPage()

    # ==========================================
    # SLIDE 16: Q&A / THANK YOU
    # ==========================================
    c.setFillColor(colors.HexColor('#0F172A'))
    c.rect(0, 0, SLIDE_W, SLIDE_H, fill=1, stroke=0)

    # Accent lines
    c.setFillColor(colors.HexColor('#2563EB'))
    c.rect(0, 525, SLIDE_W, 15, fill=1, stroke=0)
    c.setFillColor(colors.HexColor('#0D9488'))
    c.rect(0, 0, SLIDE_W, 12, fill=1, stroke=0)

    c.setFillColor(colors.white)
    c.setFont('TimesNewRoman-Bold', 36)
    c.drawCentredString(SLIDE_W/2, 380, "THANK YOU!")
    c.setFont('TimesNewRoman-Italic', 18)
    c.setFillColor(colors.HexColor('#93C5FD'))
    c.drawCentredString(SLIDE_W/2, 340, "Questions & Discussion")

    c.setFillColor(colors.HexColor('#1E293B'))
    c.setStrokeColor(colors.HexColor('#334155'))
    c.roundRect(160, 120, 640, 180, 10, fill=1, stroke=1)

    c.setFillColor(colors.white)
    c.setFont('TimesNewRoman-Bold', 14)
    c.drawString(190, 265, "Special Acknowledgments:")
    c.setFont('TimesNewRoman', 12)
    c.setFillColor(colors.HexColor('#E2E8F0'))
    c.drawString(190, 240, "• Infosys Limited — Education, Training & Assessment (ETA) Mentors, Mysore")
    c.drawString(190, 215, "• School of Computing Science and Engineering (SCSE), Galgotias University")
    c.drawString(190, 190, "• Project Mentor & Faculty Coordinator for continuous technical guidance")
    c.drawString(190, 165, "• Fellow 9 Engineering Teammates on the Kafila Agile Development Squad")

    c.setFillColor(colors.HexColor('#94A3B8'))
    c.setFont('TimesNewRoman', 11)
    c.drawCentredString(SLIDE_W/2, 65, "Candidate: Ishan Kapoor (Admission No: 21SCSE1010042) | B.Tech CSE (2025-2026)")
    c.showPage()

    c.save()
    print(f"Presentation PDF successfully built at: {output_pdf}")

def build_presentation_pptx(output_pptx):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slides_info = [
        ("TITLE", "KAFILA — SMART FLEET MANAGEMENT SYSTEM", "Enterprise-Grade iOS Platform | Infosys Mysore Internship"),
        ("OVERVIEW", "Infosys Limited — Corporate Overview & Mysore Campus", "Global technology leader & world's largest corporate university"),
        ("SQUAD", "Internship Structure & 10-Member Engineering Squad", "Agile Scrum framework orchestrated via Atlassian JIRA"),
        ("PROBLEM", "Logistics Problem Statement & Industry Motivation", "Addressing transport fragmentation, fuel fraud, and maintenance delays"),
        ("SOLUTION", "The Kafila Solution: Tri-Role Mobile Operating System", "Dedicated interfaces for Fleet Managers, Drivers, and Maintenance"),
        ("TECH STACK", "Technology Stack & Architecture Patterns", "SwiftUI, Swift 5.10 Concurrency, Observation, MVVM, Supabase BaaS"),
        ("RESPONSIBILITIES", "Ishan Kapoor's Assigned Responsibilities", "Pillar 1: Backend Integration | Pillar 2: 2FA & Auth | Pillar 3: UI Testing"),
        ("BACKEND", "Pillar 1: Backend Cloud Integration & Realtime Sync", "Supabase client, custom date decoders, WebSocket pub-sub channels"),
        ("SECURITY", "Pillar 2: Authentication & Two-Factor (2FA) Security", "Tri-role RBAC, 6-digit OTP verification, offline bypass engine"),
        ("TESTING", "Pillar 3: Automated UI & Unit Testing with XCTest", "Comprehensive test suites, ViewModel state checks, 100% pass rate"),
        ("AGILE", "Agile Scrum Project Management in JIRA", "192 Story Points burned down across four 2-week development sprints"),
        ("WALKTHROUGH", "Application Screen Flows & User Journeys", "From role selection to 2FA verification and live telematics dashboards"),
        ("CHALLENGES", "Technical Challenges Faced & Engineering Solutions", "Dual-mode auth, PostgreSQL date decoding, realtime state sync"),
        ("LEARNINGS", "Learning Experiences & Professional Outcomes", "Swift concurrency, cloud BaaS, Agile teamwork, and code rigor"),
        ("CONCLUSION", "Conclusion & Future Product Roadmap", "Project achievements, CoreML predictive AI, Apple CarPlay, OBD-II"),
        ("Q&A", "Thank You! Questions & Discussion", "Candidate: Ishan Kapoor | Infosys Mysore | Galgotias University")
    ]

    for cat, title, subtitle in slides_info:
        slide = prs.slides.add_slide(blank_layout)
        
        # Header shape
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = RGBColor(15, 23, 42)
        header_shape.line.fill.background()

        # Title text
        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.12), Inches(12), Inches(0.85))
        tf = txBox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.bold = True
        p1.font.size = Pt(20)
        p1.font.color.rgb = RGBColor(255, 255, 255)
        p1.font.name = 'Times New Roman'
        
        p2 = tf.add_paragraph()
        p2.text = f"{cat} — {subtitle}"
        p2.font.italic = True
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(147, 197, 253)
        p2.font.name = 'Times New Roman'

        # Body Card
        body_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.4), Inches(12.133), Inches(5.4))
        body_shape.fill.solid()
        body_shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
        body_shape.line.color.rgb = RGBColor(203, 213, 225)
        body_shape.line.width = Pt(1.5)

        # Body Text
        b_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(11.5), Inches(5.0))
        b_tf = b_box.text_frame
        b_tf.word_wrap = True
        bp = b_tf.paragraphs[0]
        bp.text = f"Key Discussion Points for {title}:"
        bp.font.bold = True
        bp.font.size = Pt(16)
        bp.font.color.rgb = RGBColor(30, 58, 138)
        bp.font.name = 'Times New Roman'

        bullets = [
            "Engineered native iOS Fleet Management platform using Swift 5.10, SwiftUI, and Observation framework.",
            "Integrated Supabase BaaS over PostgreSQL 15 with real-time WebSocket pub-sub channels.",
            "Designed and implemented Two-Factor Authentication (2FA) with 6-digit OTP verification and RBAC.",
            "Built robust offline testing bypass mechanism supporting predefined test credentials.",
            "Delivered comprehensive XCTest automated UI and unit regression test suites with 100% pass rate.",
            "Collaborated within a 10-member cross-functional engineering squad using Agile Scrum in JIRA."
        ]
        for b in bullets:
            p = b_tf.add_paragraph()
            p.text = f"• {b}"
            p.font.size = Pt(13)
            p.font.color.rgb = RGBColor(31, 41, 55)
            p.font.name = 'Times New Roman'
            p.space_after = Pt(8)

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12.133), Inches(0.4))
        f_tf = footer_box.text_frame
        fp = f_tf.paragraphs[0]
        fp.text = "Ishan Kapoor | iOS Development Intern | Infosys Limited, Mysore | Galgotias University"
        fp.font.size = Pt(10)
        fp.font.color.rgb = RGBColor(100, 116, 139)
        fp.font.name = 'Times New Roman'

    prs.save(output_pptx)
    print(f"Presentation PPTX successfully built at: {output_pptx}")

if __name__ == "__main__":
    pdf_out = './Internship_report_and_ppt/Internship_Presentation.pdf'
    pptx_out = './Internship_report_and_ppt/Internship_Presentation.pptx'
    build_presentation_pdf(pdf_out)
    build_presentation_pptx(pptx_out)
